"""Header-oriented .xlsx/.xlsm operations with explicit save semantics."""
from pathlib import Path
from tempfile import NamedTemporaryFile
from .errors import ConfigurationError, ValidationError


def _library():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError as exc:
        raise ConfigurationError("Instale nexus-sdk[excel] para trabalhar com planilhas.") from exc
    return Workbook, load_workbook


class Spreadsheet:
    def __init__(self, workbook, path=None, *, read_only=False):
        self.native = workbook
        self.path = Path(path) if path else None
        self.read_only = read_only

    @classmethod
    def create(cls, path, *, sheet="Dados", headers=None):
        Workbook, _ = _library()
        book = Workbook()
        book.active.title = sheet
        if headers:
            cls._validate_headers(headers)
            book.active.append(list(headers))
        return cls(book, path)

    @classmethod
    def open(cls, path, *, data_only=False):
        _, load_workbook = _library()
        target = Path(path)
        return cls(load_workbook(target, data_only=data_only, keep_vba=target.suffix.lower() == ".xlsm"), target, read_only=data_only)

    @staticmethod
    def _validate_headers(headers):
        if not headers or any(not isinstance(value, str) or not value.strip() for value in headers) or len(set(headers)) != len(headers):
            raise ValidationError("Cabeçalhos devem ser únicos e não vazios.")

    def worksheet(self, sheet="Dados"):
        if sheet not in self.native:
            raise ValidationError(f"A aba {sheet!r} não existe.")
        return self.native[sheet]

    def add_sheet(self, name, *, headers=None):
        if name in self.native:
            raise ValidationError(f"A aba {name!r} já existe.")
        if headers:
            self._validate_headers(headers)
        ws = self.native.create_sheet(name)
        if headers:
            ws.append(list(headers))
        return ws

    def headers(self, *, sheet="Dados"):
        ws = self.worksheet(sheet)
        values = [cell.value for cell in ws[1]]
        self._validate_headers(values)
        return values

    def rows(self, *, sheet="Dados"):
        ws = self.worksheet(sheet)
        headers = self.headers(sheet=sheet)
        for cells in ws.iter_rows(min_row=2, values_only=True):
            if any(value is not None for value in cells):
                yield dict(zip(headers, cells))

    def append(self, values: dict, *, sheet="Dados"):
        headers = self.headers(sheet=sheet)
        self._check_columns(values, headers)
        self.worksheet(sheet).append([values.get(name) for name in headers])

    def find(self, column, value, *, sheet="Dados") -> list[dict]:
        if column not in self.headers(sheet=sheet):
            raise ValidationError(f"Coluna {column!r} não existe.")
        return [row for row in self.rows(sheet=sheet) if row[column] == value]

    def update_row(self, row_number: int, changes: dict, *, sheet="Dados"):
        headers = self.headers(sheet=sheet)
        self._check_columns(changes, headers)
        ws = self.worksheet(sheet)
        if row_number < 2 or row_number > ws.max_row:
            raise ValidationError("Número da linha fora dos dados da planilha (cabeçalho é a linha 1).")
        for name, updated in changes.items():
            ws.cell(row_number, headers.index(name) + 1).value = updated

    @staticmethod
    def _check_columns(values, headers):
        if not isinstance(values, dict) or set(values) - set(headers):
            raise ValidationError("Informe um dicionário com colunas existentes na planilha.")

    def update_where(self, column, value, changes: dict, *, sheet="Dados", all_matches=False) -> int:
        headers = self.headers(sheet=sheet)
        if column not in headers:
            raise ValidationError(f"Coluna {column!r} não existe.")
        self._check_columns(changes, headers)
        if not changes:
            raise ValidationError("Informe ao menos uma coluna para atualizar.")
        ws = self.worksheet(sheet)
        matches = [row for row in range(2, ws.max_row + 1) if ws.cell(row, headers.index(column) + 1).value == value]
        if len(matches) > 1 and not all_matches:
            raise ValidationError("Mais de uma linha corresponde ao filtro; use all_matches=True.")
        for row in matches:
            self.update_row(row, changes, sheet=sheet)
        return len(matches)

    def save(self, path=None) -> Path:
        if self.read_only:
            raise ConfigurationError("Planilhas abertas com data_only=True não podem ser salvas, pois isso removeria as fórmulas.")
        target = Path(path) if path else self.path
        if target is None:
            raise ConfigurationError("Informe um caminho para salvar a planilha.")
        if target.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValidationError("Use um arquivo .xlsx ou .xlsm.")
        target.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(dir=target.parent, suffix=target.suffix, delete=False) as stream:
            temporary = Path(stream.name)
        try:
            self.native.save(temporary)
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        self.path = target
        return target

    def close(self):
        self.native.close()
